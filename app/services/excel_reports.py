from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

COLUMN_MAP={
    "Kundenauftragsnummer":"Kundenauftragsnummer",
    "Dossier":"Dossier",
    "Ladetermin":"Ladetermin",
    "Liefertermin":"Liefertermin",
    "Beladeadresse":"Beladeadresse",
    "Entladeadresse":"Entladeadresse",
    "Fahrzeug":"Fahrzeug",
    "Unternehmer":"Unternehmer",
}

DAYS={0:"Montag",1:"Dienstag",2:"Mittwoch",3:"Donnerstag",4:"Freitag",5:"Samstag",6:"Sonntag"}

def _load(path):
    xl=pd.ExcelFile(path)
    sheet="Disposition" if "Disposition" in xl.sheet_names else xl.sheet_names[0]
    df=pd.read_excel(path,sheet_name=sheet,header=1)
    out=pd.DataFrame()
    for s,t in COLUMN_MAP.items():
        if s in df.columns:
            out[t]=df[s]
    dt=pd.to_datetime(out["Liefertermin"],errors="coerce",dayfirst=True)
    out["Liefertermin"]=dt
    out["Kalenderwoche"]=dt.dt.isocalendar().week.astype("Int64")
    out["Jahr"]=dt.dt.year
    out["Wochentag"]=dt.dt.dayofweek.map(DAYS)
    return out

def _autosize(ws):
    for c in ws.columns:
        l=max(len(str(x.value or "")) for x in c)
        ws.column_dimensions[c[0].column_letter].width=min(max(l+2,12),45)

def create_reports(excel_path, output_root="Auswertungen"):
    df=_load(excel_path)
    if df.empty:
        raise ValueError("Keine Daten gefunden.")
    kw=int(df["Kalenderwoche"].mode().iloc[0]); jahr=int(df["Jahr"].mode().iloc[0])
    out=Path(output_root)/f"{jahr}_KW{kw:02d}"
    out.mkdir(parents=True,exist_ok=True)
    # 1 Wochenübersicht
    df.sort_values(["Liefertermin","Fahrzeug"],na_position="last").assign(
        Liefertermin=df["Liefertermin"].dt.strftime("%d.%m.%Y")
    ).to_excel(out/"01_Wochenübersicht.xlsx",index=False)
    # 2 Fuhrpark
    fuhr=df.groupby("Fahrzeug",dropna=False).size().reset_index(name="Aufträge").sort_values("Aufträge",ascending=False)
    fuhr.to_excel(out/"02_Fuhrpark.xlsx",index=False)
    # 3 Unternehmer
    unt=df.groupby("Unternehmer",dropna=False).size().reset_index(name="Aufträge").sort_values("Aufträge",ascending=False)
    unt.to_excel(out/"03_Unternehmer.xlsx",index=False)
    # 4 Tagesübersicht
    wb=Workbook(); wb.remove(wb.active)
    for day in DAYS.values():
        ws=wb.create_sheet(day); ws.append(list(df.columns))
        for cell in ws[1]: cell.font=Font(bold=True)
        part=df[df["Wochentag"]==day].copy()
        if not part.empty: part["Liefertermin"]=part["Liefertermin"].dt.strftime("%d.%m.%Y")
        for row in part.fillna("").itertuples(index=False):
            ws.append(list(row))
        _autosize(ws)
    wb.save(out/"04_Tagesübersicht.xlsx")
    # 5 Kennzahlen
    wb=Workbook(); ws=wb.active; ws.title="Kennzahlen"
    rows=[
      ("Kennzahl","Wert"),
      ("Kalenderwoche",kw),
      ("Jahr",jahr),
      ("Aufträge",len(df)),
      ("Fahrzeuge",df["Fahrzeug"].dropna().nunique()),
      ("Unternehmer",df["Unternehmer"].dropna().nunique()),
      ("Aufträge je Fahrzeug",round(len(df)/max(df["Fahrzeug"].dropna().nunique(),1),2)),
      ("Aufträge je Unternehmer",round(len(df)/max(df["Unternehmer"].dropna().nunique(),1),2)),
    ]
    for r in rows: ws.append(r)
    for c in ws[1]: c.font=Font(bold=True)
    _autosize(ws); wb.save(out/"05_Kennzahlen.xlsx")
    return out
