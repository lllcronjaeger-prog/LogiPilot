from pathlib import Path
from openpyxl import load_workbook
from app.services.shipment_service import get_week_dataframe

TEMPLATES={
 'kw':'template_kw.xlsx',
 'verkauf':'template_verkauf.xlsx',
 'wochenuebersicht':'template_wochenuebersicht.xlsx'
}

def export_with_template(week:int, kind:str, out_dir:Path, template_dir:Path):
    wb=load_workbook(template_dir/TEMPLATES[kind])
    ws=wb[wb.sheetnames[0]]
    df=get_week_dataframe(week)
    # Beispiel: KW im Kopf ersetzen
    if ws['C1'].value and 'KW' in str(ws['C1'].value):
        ws['C1']=f'KW {week}'
    start=ws.max_row+2
    if not df.empty:
        for c,col in enumerate(df.columns,1):
            ws.cell(row=start,column=c).value=col
        for r,row in enumerate(df.itertuples(index=False),start+1):
            for c,val in enumerate(row,1):
                ws.cell(row=r,column=c).value=val
    out=Path(out_dir)/f'KW{week:02d}_{kind}.xlsx'
    wb.save(out)
    return out
