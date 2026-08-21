from pathlib import Path
from openpyxl import load_workbook
from app.services.shipment_service import get_week_dataframe

TEMPLATE_MAP={
    "kw":"template_kw.xlsx",
    "verkauf":"template_verkauf.xlsx",
    "wochenuebersicht":"template_wochenuebersicht.xlsx",
}

def fill_template(week:int, kind:str, template_dir:Path, output_dir:Path):
    wb=load_workbook(template_dir/TEMPLATE_MAP[kind])
    ws=wb[wb.sheetnames[0]]
    df=get_week_dataframe(week)

    # Kalenderwoche in der ersten Zeile aktualisieren
    for cell in ws[1]:
        if cell.value and "KW" in str(cell.value):
            cell.value=f"KW {week}"

    # Rohdaten ab Zeile 100 schreiben (Vorlage bleibt unangetastet)
    start_row=100
    if not df.empty:
        for c,col in enumerate(df.columns,1):
            ws.cell(row=start_row,column=c).value=col
        for r,row in enumerate(df.itertuples(index=False),start_row+1):
            for c,val in enumerate(row,1):
                ws.cell(row=r,column=c).value=val

    output_dir=Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    out=output_dir/f"KW{week:02d}_{kind}.xlsx"
    wb.save(out)
    return out
